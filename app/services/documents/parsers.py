"""Format readers.

Each one does as little as possible: pull out text and mark what looks like a
heading. Cleaning, chunking and structuring all happen once, in the pipeline,
so behaviour stays identical across formats.
"""

from __future__ import annotations

import html
import re
from pathlib import Path

from ...domain.models import ResourceKind
from .. import textanalysis as ta
from .base import Block, ExtractedDocument, UnreadableDocument


class PdfProcessor:
    kind = ResourceKind.PDF
    extensions = (".pdf",)

    def can_handle(self, path: Path, mime_type: str = "") -> bool:
        return path.suffix.lower() in self.extensions or "pdf" in mime_type

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            from pypdf import PdfReader
            from pypdf.errors import PdfReadError
        except ImportError:
            raise UnreadableDocument("PDF support is not installed on this server.") from None

        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                # An empty user password is common and worth trying.
                try:
                    reader.decrypt("")
                except Exception:  # noqa: BLE001 - any failure means we cannot read it
                    raise UnreadableDocument(
                        "This PDF is password protected. Remove the password and upload it again."
                    ) from None
            pages = reader.pages
        except UnreadableDocument:
            raise
        except (PdfReadError, OSError, ValueError):
            raise UnreadableDocument(
                "This PDF could not be opened. It may be damaged."
            ) from None

        blocks: list[Block] = []
        for number, page in enumerate(pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:  # noqa: BLE001 - one bad page must not lose the file
                continue
            for line in text.splitlines():
                line = line.strip()
                if not line:
                    continue
                blocks.append(Block(line, page=number, is_heading=ta.is_heading(line)))

        document = ExtractedDocument(
            title=_pdf_title(reader, path),
            blocks=blocks,
            page_count=len(pages),
            kind=self.kind,
        )
        if document.is_empty:
            raise UnreadableDocument(
                "No text could be read from this PDF. If it is a scan, it needs to be "
                "run through text recognition first."
            )
        return document


def _pdf_title(reader: object, path: Path) -> str:
    try:
        metadata = getattr(reader, "metadata", None)
        title = (getattr(metadata, "title", "") or "").strip()
        if title and len(title) < 120:
            return title
    except Exception:  # noqa: BLE001 - metadata is optional
        pass
    return path.stem.replace("_", " ").replace("-", " ").strip()


class DocxProcessor:
    kind = ResourceKind.DOCX
    extensions = (".docx",)

    def can_handle(self, path: Path, mime_type: str = "") -> bool:
        return path.suffix.lower() in self.extensions or "wordprocessingml" in mime_type

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            import docx
        except ImportError:
            raise UnreadableDocument("Word support is not installed on this server.") from None

        try:
            document = docx.Document(str(path))
        except Exception:  # noqa: BLE001 - python-docx raises broadly on bad files
            raise UnreadableDocument(
                "This Word file could not be opened. It may be damaged, or saved in the "
                "older .doc format, which needs converting to .docx first."
            ) from None

        blocks: list[Block] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style else "") or ""
            is_heading = style.lower().startswith(("heading", "title", "subtitle"))
            blocks.append(Block(text, is_heading=is_heading or ta.is_heading(text)))

        # Tables often carry definitions and data worth teaching.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if len(cells) >= 2:
                    blocks.append(Block(f"{cells[0]}: {' '.join(cells[1:])}"))

        extracted = ExtractedDocument(
            title=_docx_title(document, path), blocks=blocks, page_count=0, kind=self.kind
        )
        if extracted.is_empty:
            raise UnreadableDocument("This document does not contain enough text to teach from.")
        return extracted


def _docx_title(document: object, path: Path) -> str:
    try:
        title = (document.core_properties.title or "").strip()  # type: ignore[attr-defined]
        if title and len(title) < 120:
            return title
    except Exception:  # noqa: BLE001 - properties are optional
        pass
    return path.stem.replace("_", " ").replace("-", " ").strip()


class PptxProcessor:
    kind = ResourceKind.PPTX
    extensions = (".pptx",)

    def can_handle(self, path: Path, mime_type: str = "") -> bool:
        return path.suffix.lower() in self.extensions or "presentationml" in mime_type

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            from pptx import Presentation
        except ImportError:
            raise UnreadableDocument("Slide support is not installed on this server.") from None

        try:
            presentation = Presentation(str(path))
        except Exception:  # noqa: BLE001 - python-pptx raises broadly on bad files
            raise UnreadableDocument(
                "This presentation could not be opened. It may be damaged, or saved in the "
                "older .ppt format, which needs converting to .pptx first."
            ) from None

        blocks: list[Block] = []
        slides = list(presentation.slides)
        for number, slide in enumerate(slides, start=1):
            # A slide's title is a heading by construction.
            title_shape = getattr(slide.shapes, "title", None)
            if title_shape is not None and getattr(title_shape, "text", "").strip():
                blocks.append(Block(title_shape.text.strip(), page=number, is_heading=True))
            for shape in slide.shapes:
                if shape is title_shape or not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        blocks.append(Block(text, page=number))
            notes = getattr(slide, "notes_slide", None)
            if notes is not None and getattr(notes, "notes_text_frame", None) is not None:
                note_text = notes.notes_text_frame.text.strip()
                if len(note_text) > 30:
                    blocks.append(Block(note_text, page=number))

        extracted = ExtractedDocument(
            title=path.stem.replace("_", " ").strip(),
            blocks=blocks,
            page_count=len(slides),
            kind=self.kind,
        )
        if extracted.is_empty:
            raise UnreadableDocument(
                "These slides are mostly images. Only text can be taught from, so there is "
                "not enough here yet."
            )
        return extracted


class TextProcessor:
    kind = ResourceKind.TEXT
    extensions = (".txt", ".md", ".markdown", ".text", ".rst")

    def can_handle(self, path: Path, mime_type: str = "") -> bool:
        return path.suffix.lower() in self.extensions or mime_type.startswith("text/")

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            raw = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            try:
                raw = path.read_text(encoding="latin-1")
            except OSError:
                raise UnreadableDocument("This file could not be read as text.") from None
        except OSError:
            raise UnreadableDocument("This file could not be read.") from None

        is_markdown = path.suffix.lower() in (".md", ".markdown")
        blocks: list[Block] = []
        for line in raw.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            heading = False
            if is_markdown and stripped.startswith("#"):
                stripped = stripped.lstrip("#").strip()
                heading = True
            elif is_markdown and re.fullmatch(r"[-=]{3,}", stripped) and blocks:
                blocks[-1].is_heading = True      # setext underline
                continue
            else:
                heading = ta.is_heading(stripped)
            if is_markdown:
                stripped = _strip_markdown(stripped)
            if stripped:
                blocks.append(Block(stripped, is_heading=heading))

        extracted = ExtractedDocument(
            title=path.stem.replace("_", " ").replace("-", " ").strip(),
            blocks=blocks,
            kind=ResourceKind.MARKDOWN if is_markdown else ResourceKind.TEXT,
        )
        if extracted.is_empty:
            raise UnreadableDocument("There is not enough text here to build a lesson from.")
        return extracted


def _strip_markdown(text: str) -> str:
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)          # images
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)      # links
    text = re.sub(r"[*_]{1,3}([^*_]+)[*_]{1,3}", r"\1", text)  # emphasis
    text = re.sub(r"`{1,3}([^`]*)`{1,3}", r"\1", text)         # code
    text = re.sub(r"^\s*[-*+]\s+", "", text)                   # bullets
    text = re.sub(r"^\s*>\s?", "", text)                       # quotes
    return text.strip()


class NoteProcessor:
    """Text typed straight into the app rather than uploaded."""

    kind = ResourceKind.NOTE
    extensions = (".note",)

    def can_handle(self, path: Path, mime_type: str = "") -> bool:
        return path.suffix.lower() in self.extensions

    def extract(self, path: Path) -> ExtractedDocument:
        document = TextProcessor().extract(path)
        document.kind = self.kind
        return document


class HtmlProcessor:
    """A saved web page, or one fetched from a URL."""

    kind = ResourceKind.URL
    extensions = (".html", ".htm")

    def can_handle(self, path: Path, mime_type: str = "") -> bool:
        return path.suffix.lower() in self.extensions or "html" in mime_type

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            raw = path.read_text(encoding="utf-8", errors="replace")
        except OSError:
            raise UnreadableDocument("This page could not be read.") from None
        return self.from_html(raw, title=path.stem)

    @staticmethod
    def from_html(raw: str, *, title: str = "") -> ExtractedDocument:
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(raw, "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript", "form"]):
                tag.decompose()
            page_title = (soup.title.string or "").strip() if soup.title else ""
            blocks: list[Block] = []
            for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "dd", "dt", "blockquote"]):
                text = " ".join(element.get_text(" ", strip=True).split())
                if len(text) < 3:
                    continue
                blocks.append(Block(text, is_heading=element.name in ("h1", "h2", "h3", "h4")))
        except ImportError:
            # No BeautifulSoup: strip tags with a conservative regex instead.
            page_title = ""
            match = re.search(r"<title[^>]*>(.*?)</title>", raw, re.I | re.S)
            if match:
                page_title = html.unescape(match.group(1)).strip()
            body = re.sub(r"(?is)<(script|style|nav|footer|header|aside)[^>]*>.*?</\1>", " ", raw)
            body = re.sub(r"(?i)<(h[1-4])[^>]*>(.*?)</\1>", r"\n## \2\n", body)
            body = re.sub(r"(?i)</(p|div|li|tr|br)\s*/?>", "\n", body)
            body = re.sub(r"<[^>]+>", " ", body)
            blocks = []
            for line in html.unescape(body).splitlines():
                stripped = " ".join(line.split())
                if len(stripped) < 3:
                    continue
                heading = stripped.startswith("##")
                blocks.append(Block(stripped.lstrip("# ").strip(), is_heading=heading))

        extracted = ExtractedDocument(
            title=(page_title or title or "Web page")[:120],
            blocks=blocks,
            kind=ResourceKind.URL,
        )
        if extracted.is_empty:
            raise UnreadableDocument(
                "Not enough readable text was found on that page. Some sites load their "
                "content with scripts, which cannot be read this way."
            )
        return extracted


#: Tried in order. The first that says it can handle the file wins.
PROCESSORS: tuple[object, ...] = (
    PdfProcessor(),
    DocxProcessor(),
    PptxProcessor(),
    HtmlProcessor(),
    NoteProcessor(),
    TextProcessor(),
)
